"""Regenerate src/data/projects.js from "Portfolio List.xlsx".

Usage (from the portfolio-search folder):
    python scripts/generate_projects.py

To add more projects, add the sheet's Brand Name to INCLUDE
(value = the name to display on the site). Requires: pip install openpyxl

Only three columns are used: Brand Name, Industry and tags.
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT.parent / "Portfolio List.xlsx"
OUTPUT = ROOT / "src" / "data" / "projects.js"

INCLUDE = {
    "Wattle Physio": "Wattle Physio",
    "Via Wind": "Via Wind",
    "StrateAura": "StrateAura",
    "Arani": "Arani",
    "Infiled": "INFiLED",
}

def split_tags(raw):
    seen, tags = set(), []
    for tag in re.split(r"[,\n]", raw or ""):
        tag = " ".join(tag.split())
        if tag and tag.lower() not in seen:
            seen.add(tag.lower())
            tags.append(tag)
    return tags


def main():
    ws = openpyxl.load_workbook(SHEET).worksheets[0]
    header = [cell.value for cell in ws[1]]
    projects = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] not in INCLUDE:
            continue
        data = dict(zip(header, row))
        name = INCLUDE[row[0]]
        projects.append({
            "id": re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-"),
            "name": name,
            "industry": data["Industry"],
            "tags": split_tags(data["tags"]),
        })

    js = '// Generated from "Portfolio List.xlsx" by scripts/generate_projects.py — do not edit by hand.\n'
    js += "export const projects = " + json.dumps(projects, indent=2, ensure_ascii=False) + "\n"
    OUTPUT.write_text(js, encoding="utf-8")
    print(f"Wrote {len(projects)} projects to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
